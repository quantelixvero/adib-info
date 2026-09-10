# -*- coding: utf-8 -*-
import os
import glob
import re
import json
import unicodedata
import csv
from pypdf import PdfReader
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Student Directory Data from Screenshots
DIRECTORY_METADATA = {
    '139': {'room': 'Room 325', 'section': 'Section A', 'display_name': 'Abdullah Al Adib', 'group': 'Science'},
    '161': {'room': 'Room 317', 'section': 'Section B', 'display_name': 'Naimul Islam (Shanto)', 'group': 'Science'},
    '164': {'room': 'Room 317', 'section': 'Section B', 'display_name': 'Md Zahidul Islam', 'group': 'Science'},
    '326': {'room': 'Room 325', 'section': 'Section C', 'display_name': 'Md. Shakil Sheikh', 'group': 'Science'},
    '008': {'room': 'Room 318', 'section': 'Section A', 'display_name': 'Shafkat Rahman', 'group': 'Science'},
    '022': {'room': 'Room 315', 'section': 'Section A', 'display_name': 'Shahriar Al Sakib Munshi', 'group': 'Science'},
    '122': {'room': 'Room 316', 'section': 'Section A', 'display_name': 'Md. Mushfiqur Rahman', 'group': 'Science'},
    '870': {'room': 'Room 322', 'section': 'Section F', 'display_name': 'Taskin Rahman Esty', 'group': 'Science'},
    '141': {'room': 'Room 316', 'section': 'Section A', 'display_name': 'Nafis Alam Tarif', 'group': 'Science'},
    '882': {'room': 'Room 314', 'section': 'Section F', 'display_name': 'Riasadul Islam Kayes', 'group': 'Science'},
    '223': {'room': 'Room 328', 'section': 'Section B', 'display_name': 'Md. Jahidul Islam', 'group': 'Science'},
    '195': {'room': 'Room 319', 'section': 'Section B', 'display_name': 'Siam Hasan (Nator)', 'group': 'Science'},
    '881': {'room': 'Room 314', 'section': 'Section F', 'display_name': 'Md. Siam Ali (Kustia)', 'group': 'Science'},
    '322': {'room': 'Room 326', 'section': 'Section C', 'display_name': 'Saimun Islam', 'group': 'Science'},
    '863': {'room': 'Room 323', 'section': 'Section F', 'display_name': 'Ratul Hassan Pranto', 'group': 'Science'},
    '639': {'room': 'Room 323', 'section': 'Section E', 'display_name': 'Nayeem Sarwar', 'group': 'Science'},
    '140': {'room': 'Room 322', 'section': 'Section -', 'display_name': 'Md. Alauddin Mollah', 'group': 'Humanities'},
    '872': {'room': 'Room 321', 'section': 'Section F', 'display_name': 'Md. Rafiq Bin Nizam', 'group': 'Science'},
    '596': {'room': 'Room 318', 'section': 'Section D', 'display_name': 'Tashfiqur Rahman Riyad', 'group': 'Science'},
    '332': {'room': 'Room 320', 'section': 'Section C', 'display_name': 'Md. Chisty Islam Mosih', 'group': 'Science'},
    '200': {'room': 'Room 320', 'section': 'Section B', 'display_name': 'Md. Ahanaf Abrar', 'group': 'Science'},
    '100': {'room': 'Room 324', 'section': 'Section -', 'display_name': 'M M Tanvir (Rudro)', 'group': 'Business Studies'},
    '337': {'room': 'Room 327', 'section': 'Section C', 'display_name': 'Sobhan Mahmud Khan (Riasad)', 'group': 'Science'},
    '213': {'room': 'Room 328', 'section': 'Section B', 'display_name': 'Mahbubur Rahman Musa', 'group': 'Science'},
    '050': {'room': 'Room 324', 'section': 'Section A', 'display_name': 'Ariful Islam', 'group': 'Science'},
    '098': {'room': 'Room 315', 'section': 'Section -', 'display_name': 'Md Rezaul Karim', 'group': 'Humanities'},
    '171': {'room': 'Room 327', 'section': 'Section B', 'display_name': 'Muhammad Rayhanoor Sabit', 'group': 'Science'},
    '153': {'room': 'Room 326', 'section': 'Section B', 'display_name': 'Alimuzzamann Alif', 'group': 'Science'},
}

# Accurate Bengali Names for Student, Father, and Mother
BANGLA_NAMES = {
    '139': ('আবদুল্লাহ আল আদিব', 'মোঃ আকবর আলী', 'রোজিনা আকতার'),
    '161': ('নাইমুল ইসলাম (শান্ত)', 'নজরুল ইসলাম', 'মরিয়ম বেগম'),
    '164': ('মোঃ জাহিদুল ইসলাম', 'মোঃ জাহিরুল ইসলাম', 'মেরিনা আক্তার আঁখি'),
    '326': ('মোঃ শাকিল শেখ', 'মোঃ আলম শেখ', 'আসমা বেগম'),
    '008': ('শাফকাত রহমান', 'মোঃ শাহজাহান আলী', 'শারমিন সুলতানা'),
    '022': ('শাহরিয়ার আল সাকিব মুন্সী', 'শাহনেওয়াজ মুন্সী', 'মোসাম্মৎ শিরিনা আক্তার'),
    '122': ('মোঃ মুশফিকুর রহমান মহিন', 'মোঃ কামাল হোসেন', 'লুৎফুন নাহার'),
    '870': ('তাসকিন রহমান এস্টি', 'মুহাম্মদ সিদ্দিকুর রহমান', 'আঞ্জুমান আরা আরজু'),
    '141': ('মোঃ নাফিস আলম তারিফ', 'মোঃ নূহেদ আলম', 'মোছাঃ নিলুফা আইরিন'),
    '882': ('রিয়াসাদুল ইসলাম কায়েস', 'সাইদুল কাদের', 'রেশমিন আক্তার'),
    '223': ('মোঃ জাহিদুল ইসলাম', 'খোরশেদ আলম', 'জোৎস্না বেগম'),
    '195': ('সিয়াম হাসান', 'মোঃ সাজেদুর রহমান', 'মোছাঃ অনামিকা খাতুন'),
    '881': ('মোঃ সিয়াম আলী', 'মোঃ রবিউল ইসলাম', 'মোছাঃ সুলতানা খাতুন'),
    '322': ('মোঃ সাইমুন ইসলাম', 'মোঃ আবু সাঈদ', 'মোছাঃ সাবিনা ইয়াসমিন'),
    '863': ('মোঃ রাতুল হাসান প্রান্ত', 'মোঃ আলতাফ হোসেন', 'মৃত নাজমুন নাহার'),
    '639': ('নাঈম সরওয়ার', 'আবু সাঈদ মোঃ গোলাম সরওয়ার', 'সামসুননাহার'),
    '140': ('মোঃ আলাউদ্দিন মোল্লা', 'মোঃ ইসাহাক মোল্লা', 'সুফিয়া'),
    '872': ('মোঃ রফিক বিন নিজাম', 'মোঃ জাফিরুল ইসলাম', 'মোছাঃ রাবিয়া বসরী'),
    '596': ('তাশফিকুর রহমান রিয়াদ', 'মাসুম মিয়া', 'তাসমিন আক্তার'),
    '332': ('মোঃ চিশতী ইসলাম মসীহ', 'মোঃ বজলুর রহমান', 'মোছাঃ আঁখি বেগম'),
    '200': ('মোঃ আহনাফ আবরার', 'মোঃ জাহাঙ্গীর হোসেন', 'মোছাঃ জলি খাতুন'),
    '100': ('এম এম তানভীর (রুদ্র)', 'মোঃ রফিকুল ইসলাম', 'মোছাঃ তাসলিমা খাতুন সুফিয়া'),
    '337': ('সোবহান মাহমুদ খান (রিয়াসাদ)', 'মাহাবুব খাঁন', 'মোসাঃ লুৎফুন্নাহার শিমু'),
    '213': ('মাহবুবুর রহমান মুসা', 'মোঃ মিজানুর রহমান', 'জেসমিন আক্তার'),
    '050': ('মোঃ আরিফুল ইসলাম', 'মোঃ আজিজুল ইসলাম', 'মোসাঃ নাজমা বেগম'),
    '098': ('মোঃ রেজাউল করিম', 'তোফাজ্জল হোসেন', 'মোসাঃ রানু বেগম'),
    '171': ('মুহাম্মদ রায়হানুর সাবিত', 'মোঃ আবুল কাশেম', 'মোছাঃ সালমা বেগম'),
    '153': ('মোঃ আলিমুজ্জামান আলিফ', 'মোঃ নুর জামান মিয়া', 'মোছাঃ আয়েশা সিদ্দিকা'),
}

# Website students.json metadata
site_students_map = {}
json_path = r'c:\Users\USER\Downloads\website DChall\hall-main\data\students.json'
if os.path.exists(json_path):
    with open(json_path, encoding='utf-8') as f:
        sj = json.load(f)
        for s in sj.get('students', []):
            roll = s.get('short_roll', '').zfill(3)
            site_students_map[roll] = s

def clean_text(t):
    if not t:
        return ''
    t = unicodedata.normalize('NFKD', t)
    t = t.replace('\ufb00', 'ff').replace('\ufb01', 'fi').replace('\ufb02', 'fl').replace('\ufb03', 'ffi').replace('\ufb04', 'ffl')
    t = re.sub(r'[ \t]+', ' ', t)
    return t.strip()

def parse_student(app_pdf_path):
    fn = os.path.basename(app_pdf_path)
    m_roll = re.search(r'-(\d{3})\.pdf$', fn)
    short_roll = m_roll.group(1) if m_roll else ''
    
    text = clean_text(PdfReader(app_pdf_path).pages[0].extract_text())
    
    dir_info = DIRECTORY_METADATA.get(short_roll, {})
    site_info = site_students_map.get(short_roll, {})
    bn_info = BANGLA_NAMES.get(short_roll, ('', '', ''))
    
    # Class roll & admission roll
    m_croll = re.search(r'Class Roll:\s*(\d+)', text)
    full_class_roll = m_croll.group(1).strip() if m_croll else ''
    
    m_aroll = re.search(r'Admission Roll:\s*(\d+)', text)
    admission_roll = m_aroll.group(1).strip() if m_aroll else ''
    
    # Group & Session
    m_grp = re.search(r'Group:\s*(.*?)\n', text)
    group_name = m_grp.group(1).strip() if m_grp else dir_info.get('group', 'Science')
    if not group_name.startswith('HSC -'):
        group_name = f"HSC - {group_name}"
        
    m_sess = re.search(r'Session:\s*([0-9\-]+)', text)
    session = m_sess.group(1).strip() if m_sess else '2025-2026'
    
    m_trk = re.search(r'Track No:\s*(\d+)', text)
    track_no = m_trk.group(1).strip() if m_trk else ''
    
    m_eiin = re.search(r'EIIN:\s*(\d+)', text)
    eiin = m_eiin.group(1).strip() if m_eiin else '107977'
    
    m_trx = re.search(r'Transaction No:\s*([A-Za-z0-9]+)', text)
    transaction_no = m_trx.group(1).strip() if m_trx else ''
    
    m_pmode = re.search(r'Payment Mode:\s*([A-Za-z0-9]+)', text)
    payment_mode = m_pmode.group(1).strip() if m_pmode else ''
    
    m_date = re.search(r'Application Date:\s*([0-9A-Za-z\-]+)', text)
    app_date = m_date.group(1).strip() if m_date else ''
    
    # Student Name (English)
    m_sname = re.search(r'01\.\s*Student Name\s*\(English\)\s*:\s*(.*?)\s*(?:\([^\)]*বাংলা|$)', text)
    student_name_en = m_sname.group(1).strip() if m_sname else dir_info.get('display_name', '')
    
    student_name_bn = bn_info[0]
    
    # Student NID / Birth Reg
    m_snid = re.search(r"Student's NID/Birth Reg\.\s*:\s*(\d+)", text)
    student_nid = m_snid.group(1).strip() if m_snid else ''
    
    # Gender
    m_gen = re.search(r"Gender\s*:\s*([A-Za-z]+)", text)
    gender = m_gen.group(1).strip() if m_gen else 'Male'
    
    # Student Phone & Email
    m_sphone = re.search(r"Student's Phone\s*:\s*([0-9]+)", text)
    student_phone = m_sphone.group(1).strip() if m_sphone else ''
    
    m_email = re.search(r"Student's E-mail\s*:\s*([^\s\n]+@[^\s\n]+)", text)
    student_email = m_email.group(1).strip() if m_email else site_info.get('email', '')
    
    # Father's Name (English & Bangla)
    m_fname = re.search(r"02\.\s*Father's Name\s*:\s*(.*?)\s*(?:\([^\)]*বাংলা|$)", text)
    father_name_en = m_fname.group(1).strip() if m_fname else ''
    father_name_bn = bn_info[1]
    
    # Father NID & Phone
    m_fnid = re.search(r"Father's/Guardian's NID\s*:\s*(\d+)", text)
    father_nid = m_fnid.group(1).strip() if m_fnid else ''
    
    m_fphone = re.search(r"Father's/Guardian's Phone\s*:\s*([0-9]+)", text)
    father_phone = m_fphone.group(1).strip() if m_fphone else ''
    
    # Mother's Name (English & Bangla)
    m_mname = re.search(r"03\.\s*Mother's Name\s*:\s*(.*?)\s*(?:\([^\)]*বাংলা|$)", text)
    mother_name_en = m_mname.group(1).strip() if m_mname else ''
    mother_name_bn = bn_info[2]
    
    # Mother NID & Phone
    m_mnid = re.search(r"Mother's NID\s*:\s*(\d+)", text)
    mother_nid = m_mnid.group(1).strip() if m_mnid else ''
    
    m_mphone = re.search(r"Mother's Phone\s*:\s*([0-9]+)", text)
    mother_phone = m_mphone.group(1).strip() if m_mphone else ''
    
    # Permanent & Present Address
    m_perm = re.search(r"04\.\s*Permanent Address\s*:\s*(.*?)(?=05\.\s*Present Address)", text, re.DOTALL)
    perm_addr = ' '.join(m_perm.group(1).split()).strip() if m_perm else ''
    
    m_pres = re.search(r"05\.\s*Present Address\s*:\s*(.*?)(?=06\.\s*Local Guardian)", text, re.DOTALL)
    pres_addr = ' '.join(m_pres.group(1).split()).strip() if m_pres else ''
    
    # District
    m_dist = re.search(r"District:\s*([A-Za-z]+)", perm_addr)
    if not m_dist:
        m_dist = re.search(r"District:\s*([A-Za-z]+)", text)
    district = m_dist.group(1).strip() if m_dist else ''
    
    # Local Guardian
    m_guard = re.search(r"06\.\s*Local Guardian Name, Address & Phone\s*:\s*(.*?)(?=07\.\s*Nationality)", text, re.DOTALL)
    local_guardian = ' '.join(m_guard.group(1).split()).strip() if m_guard else ''
    
    nationality = 'Bangladeshi'
    
    # Father's Occupation & Annual Income
    m_focc = re.search(r"08\.\s*Father's Occupation\s*:\s*(.*?)\s*11\.\s*Father's Annual Income\s*:\s*([0-9]+)", text)
    father_occ = m_focc.group(1).strip() if m_focc else ''
    father_income = m_focc.group(2).strip() if m_focc else ''
    
    # Date of birth & Quota
    m_dob = re.search(r"09\.\s*Date of Birth\s*:\s*([0-9A-Za-z\-]+)\s*12\.\s*Quota\s*:\s*(.*?)(?=\n|10\.)", text)
    dob = m_dob.group(1).strip() if m_dob else ''
    quota = m_dob.group(2).strip() if m_dob else ''
    
    # Religion & Blood Group
    m_rel = re.search(r"10\.\s*Religion\s*:\s*(.*?)\s*13\.\s*Blood Group\s*:\s*([A-Za-z\+\-]+)", text)
    religion = m_rel.group(1).strip() if m_rel else 'Islam'
    blood_group = m_rel.group(2).strip() if m_rel else ''
    
    # SSC Exam Details
    m_ssc = re.search(r'SSC\s+(\d+)\s+(\d+)\s+([0-9\-]+)\s+(.*?)\s+(\d{4})\s+([A-Za-z]+)\s+([0-9\.]+)', text)
    if m_ssc:
        ssc_roll = m_ssc.group(1).strip()
        ssc_reg_no = m_ssc.group(2).strip()
        ssc_session = m_ssc.group(3).strip()
        ssc_division = m_ssc.group(4).strip()
        ssc_year = m_ssc.group(5).strip()
        ssc_board = m_ssc.group(6).strip()
        ssc_gpa = m_ssc.group(7).strip()
    else:
        ssc_roll = admission_roll
        ssc_reg_no = ''
        ssc_session = ''
        ssc_division = ''
        ssc_year = '2025'
        ssc_board = ''
        ssc_gpa = ''
        
    # Subjects: Elective and Fourth
    m_elec = re.search(r'([0-9]+ - [A-Za-z ]+?)\s*(?:[0-9]+ - [A-Za-z ]+?)?\s*elective', text, re.IGNORECASE)
    elective_subj = m_elec.group(1).strip() if m_elec else ''
    
    m_fourth = re.search(r'([0-9]+ - [A-Za-z ]+?)\s*(?:[0-9]+ - [A-Za-z ]+?)?\s*fourth', text, re.IGNORECASE)
    fourth_subj = m_fourth.group(1).strip() if m_fourth else ''
    
    # Admission Fee
    m_fee = re.search(r'Admission Fee:\s*(.*?)(?=\n|Student\'s Signature)', text)
    admission_fee = m_fee.group(1).strip() if m_fee else 'TK.4710/='
    # Clean fee number
    m_amt = re.search(r'(\d+)', admission_fee)
    fee_amount = m_amt.group(1) if m_amt else '4710'
    
    # Cross reference Payment Receipt
    receipt_files = glob.glob(f'Payment Receipts/*{short_roll}.pdf')
    receipt_fn = ''
    pay_mfs = payment_mode
    pay_trx = transaction_no
    pay_date = app_date
    pay_type = 'New Admission without ICT'
    pay_amt = fee_amount
    
    if receipt_files:
        receipt_fn = os.path.basename(receipt_files[0])
        try:
            r_text = clean_text(PdfReader(receipt_files[0]).pages[0].extract_text())
            m_pay = re.search(r'(bKash|rocket|nagad|[A-Za-z]+)\s+(\S+)\s+(\S+)\s+(.*?)\s+([\d,]+)', r_text)
            if m_pay:
                pay_mfs = m_pay.group(1).strip()
                pay_trx = m_pay.group(2).strip()
                pay_date = m_pay.group(3).strip()
                pay_type = m_pay.group(4).strip()
                pay_amt = m_pay.group(5).strip().replace(',', '')
        except Exception:
            pass

    fb_url = site_info.get('facebook', '')
    if fb_url == 'https://facebook.com/' or not fb_url.startswith('http'):
        fb_url = ''

    return {
        'short_roll': short_roll,
        'full_class_roll': full_class_roll,
        'admission_roll': admission_roll,
        'hall_room': dir_info.get('room', ''),
        'group': group_name,
        'section': dir_info.get('section', ''),
        'student_name_en': student_name_en,
        'student_name_bn': student_name_bn,
        'student_phone': student_phone,
        'student_nid': student_nid,
        'student_email': student_email,
        'gender': gender,
        'date_of_birth': dob,
        'blood_group': blood_group,
        'religion': religion,
        'quota': quota,
        'nationality': nationality,
        'father_name_en': father_name_en,
        'father_name_bn': father_name_bn,
        'father_phone': father_phone,
        'father_nid': father_nid,
        'father_occupation': father_occ,
        'father_annual_income': father_income,
        'mother_name_en': mother_name_en,
        'mother_name_bn': mother_name_bn,
        'mother_phone': mother_phone,
        'mother_nid': mother_nid,
        'local_guardian': local_guardian,
        'present_address': pres_addr,
        'permanent_address': perm_addr,
        'district': district,
        'ssc_roll': ssc_roll,
        'ssc_reg_no': ssc_reg_no,
        'ssc_board': ssc_board,
        'ssc_passing_year': ssc_year,
        'ssc_gpa': ssc_gpa,
        'ssc_session': ssc_session,
        'ssc_division': ssc_division,
        'elective_subject': elective_subj,
        'fourth_subject': fourth_subj,
        'admission_session': session,
        'track_no': track_no,
        'eiin': eiin,
        'payment_mfs': pay_mfs,
        'transaction_id': pay_trx,
        'payment_date': pay_date,
        'fee_type': pay_type,
        'fee_amount_bdt': pay_amt,
        'hall_name': 'International Hall',
        'facebook_url': fb_url,
        'app_form_file': fn,
        'payment_receipt_file': receipt_fn
    }

# Read and parse all PDFs
app_files = sorted(glob.glob('Application Form/*.pdf'))
records = [parse_student(f) for f in app_files]

# Sort by short roll integer or class roll
records.sort(key=lambda x: int(x['short_roll']) if x['short_roll'].isdigit() else 9999)

print(f"Total parsed records: {len(records)}")

# Columns definition: (Field Key, Header Name, Sample Width)
COLUMNS = [
    ('sl_no', 'Sl No.', 8),
    ('hall_room', 'Room No.', 12),
    ('short_roll', 'College Roll (Short)', 18),
    ('full_class_roll', 'College Roll (Full)', 22),
    ('student_name_en', 'Student Name (English)', 28),
    ('student_name_bn', 'Student Name (বাংলা)', 26),
    ('group', 'Academic Group', 18),
    ('section', 'Section', 12),
    ('student_phone', 'Student Mobile', 18),
    ('student_nid', 'Student NID / Birth Reg', 24),
    ('date_of_birth', 'Date of Birth', 15),
    ('blood_group', 'Blood Group', 14),
    ('gender', 'Gender', 10),
    ('religion', 'Religion', 12),
    ('student_email', 'Student Email', 26),
    ('father_name_en', "Father's Name (English)", 26),
    ('father_name_bn', "Father's Name (বাংলা)", 26),
    ('father_phone', "Father's Mobile", 18),
    ('father_nid', "Father's NID", 22),
    ('father_occupation', "Father's Occupation", 20),
    ('father_annual_income', "Father's Annual Income (BDT)", 24),
    ('mother_name_en', "Mother's Name (English)", 26),
    ('mother_name_bn', "Mother's Name (বাংলা)", 26),
    ('mother_phone', "Mother's Mobile", 18),
    ('mother_nid', "Mother's NID", 22),
    ('local_guardian', 'Local Guardian Name & Phone', 32),
    ('present_address', 'Present Address', 40),
    ('permanent_address', 'Permanent Address', 42),
    ('district', 'District', 16),
    ('ssc_roll', 'SSC Exam Roll', 16),
    ('ssc_reg_no', 'SSC Registration No', 22),
    ('ssc_board', 'SSC Board', 16),
    ('ssc_passing_year', 'SSC Year', 12),
    ('ssc_gpa', 'SSC GPA', 12),
    ('ssc_session', 'SSC Session', 14),
    ('ssc_division', 'SSC Division', 16),
    ('admission_roll', 'Admission Roll', 16),
    ('elective_subject', 'Elective Subject', 24),
    ('fourth_subject', '4th Subject (Optional)', 24),
    ('admission_session', 'Admission Session', 18),
    ('track_no', 'Track No', 14),
    ('eiin', 'EIIN', 12),
    ('payment_mfs', 'Payment Method (MFS)', 20),
    ('transaction_id', 'Transaction ID (TrxID)', 22),
    ('payment_date', 'Payment Date', 16),
    ('fee_type', 'Fee Type', 26),
    ('fee_amount_bdt', 'Admission Fee (BDT)', 20),
    ('hall_name', 'Hostel / Hall Name', 20),
    ('facebook_url', 'Facebook Profile', 30),
    ('app_form_file', 'Application Form PDF', 26),
    ('payment_receipt_file', 'Payment Receipt PDF', 26),
]

# 1. GENERATE EXCEL WORKBOOK (.xlsx)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Students Information"
ws.views.sheetView[0].showGridLines = True

# Title banner
ws.merge_cells('A1:AY1')
ws['A1'] = "DHAKA COLLEGE - INTERNATIONAL HALL STUDENT DIRECTORY & ADMISSION RECORDS"
ws['A1'].font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws['A1'].fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
ws.row_dimensions[1].height = 40

# Subtitle banner
ws.merge_cells('A2:AY2')
ws['A2'] = "Session: 2025-2026 | Total Boarders: 28 Students | Comprehensive Extracted Data (Roll, Reg, College Roll, Parents, Phone, NID, Hall Room, Fees)"
ws['A2'].font = Font(name='Calibri', size=11, italic=True, color='E0E0E0')
ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws['A2'].fill = PatternFill(start_color='2C4A75', end_color='2C4A75', fill_type='solid')
ws.row_dimensions[2].height = 24

# Header row
ws.row_dimensions[3].height = 32
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='0F4C81', end_color='0F4C81', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for col_idx, (key, title, width) in enumerate(COLUMNS, 1):
    cell = ws.cell(row=3, column=col_idx, value=title)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = max(width, len(title) + 4)

# Data rows
data_font = Font(name='Calibri', size=10)
even_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
odd_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

text_cols = {
    'short_roll', 'full_class_roll', 'student_phone', 'student_nid',
    'father_phone', 'father_nid', 'mother_phone', 'mother_nid',
    'ssc_roll', 'ssc_reg_no', 'admission_roll', 'track_no', 'eiin', 'transaction_id'
}

center_cols = {
    'sl_no', 'hall_room', 'short_roll', 'full_class_roll', 'section',
    'student_phone', 'student_nid', 'date_of_birth', 'blood_group', 'gender',
    'father_phone', 'father_nid', 'mother_phone', 'mother_nid',
    'ssc_roll', 'ssc_reg_no', 'ssc_passing_year', 'ssc_gpa', 'ssc_session',
    'admission_roll', 'admission_session', 'track_no', 'eiin',
    'payment_mfs', 'transaction_id', 'payment_date', 'fee_amount_bdt'
}

for row_idx, rec in enumerate(records, 4):
    ws.row_dimensions[row_idx].height = 22
    row_fill = odd_fill if row_idx % 2 == 0 else even_fill
    rec['sl_no'] = str(row_idx - 3)
    
    for col_idx, (key, title, width) in enumerate(COLUMNS, 1):
        val = rec.get(key, '')
        cell = ws.cell(row=row_idx, column=col_idx)
        
        # Explicit text format to preserve leading zeros in phones & rolls
        if key in text_cols:
            cell.number_format = '@'
            cell.value = str(val) if val else ''
        elif key == 'sl_no':
            cell.value = int(val)
        elif key == 'fee_amount_bdt' and val.isdigit():
            cell.value = int(val)
            cell.number_format = '#,##0'
        elif key == 'father_annual_income' and val.isdigit():
            cell.value = int(val)
            cell.number_format = '#,##0'
        else:
            cell.value = val
            
        cell.font = data_font
        cell.fill = row_fill
        cell.border = thin_border
        
        if key in center_cols:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Freeze panes under header
ws.freeze_panes = 'E4'

excel_path = r'c:\Users\USER\Downloads\website DChall\Pdf for students\Dhaka_College_Students_Full_Info.xlsx'
wb.save(excel_path)
print(f"Saved Excel workbook to {excel_path}")

# 2. GENERATE CSV (.csv) WITH UTF-8 BOM FOR GOOGLE SHEETS & EXCEL
csv_path = r'c:\Users\USER\Downloads\website DChall\Pdf for students\Dhaka_College_Students_Full_Info.csv'
with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.writer(f)
    # Write header
    writer.writerow([title for key, title, width in COLUMNS])
    # Write data rows
    for idx, rec in enumerate(records, 1):
        rec['sl_no'] = str(idx)
        row = []
        for key, title, width in COLUMNS:
            val = str(rec.get(key, ''))
            # For phones and NIDs in CSV, enclose with tab or apostrophe if needed or plain
            row.append(val)
        writer.writerow(row)

print(f"Saved CSV file to {csv_path}")

# 3. GENERATE INTERACTIVE HTML FILE
html_path = r'c:\Users\USER\Downloads\website DChall\Pdf for students\Dhaka_College_Students_Full_Info.html'
html_content = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Dhaka College - International Hall Student Directory</title>
<style>
  body { font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; margin: 0; padding: 20px; background: #f0f2f5; color: #1c1e21; }
  .header-card { background: linear-gradient(135deg, #1B365D, #0F4C81); color: white; padding: 24px 30px; border-radius: 12px; margin-bottom: 20px; box-shadow: 0 4px 12px rgba(0,0,0,0.15); }
  .header-card h1 { margin: 0 0 8px 0; font-size: 24px; font-weight: 700; }
  .header-card p { margin: 0; opacity: 0.9; font-size: 14px; }
  .toolbar { display: flex; gap: 12px; margin-bottom: 16px; flex-wrap: wrap; align-items: center; }
  .search-box { flex: 1; min-width: 250px; padding: 10px 14px; border: 1px solid #ced4da; border-radius: 8px; font-size: 14px; outline: none; }
  .search-box:focus { border-color: #0F4C81; box-shadow: 0 0 0 3px rgba(15,76,129,0.2); }
  .btn { padding: 10px 18px; border-radius: 8px; border: none; font-weight: 600; font-size: 13px; cursor: pointer; display: inline-flex; align-items: center; gap: 6px; text-decoration: none; }
  .btn-primary { background: #0F4C81; color: white; }
  .btn-primary:hover { background: #1B365D; }
  .btn-success { background: #28a745; color: white; }
  .btn-success:hover { background: #218838; }
  .stats-bar { display: flex; gap: 16px; margin-bottom: 16px; flex-wrap: wrap; }
  .stat-chip { background: white; padding: 10px 16px; border-radius: 8px; font-size: 13px; font-weight: 500; box-shadow: 0 2px 4px rgba(0,0,0,0.05); }
  .stat-chip span { font-weight: 700; color: #0F4C81; }
  .table-container { background: white; border-radius: 12px; overflow-x: auto; box-shadow: 0 4px 12px rgba(0,0,0,0.08); max-height: 75vh; }
  table { border-collapse: collapse; width: 100%; font-size: 12px; white-space: nowrap; }
  th { background: #1B365D; color: white; padding: 12px 14px; position: sticky; top: 0; z-index: 10; font-weight: 600; text-align: left; }
  td { padding: 10px 14px; border-bottom: 1px solid #edf2f7; }
  tr:nth-child(even) { background: #f8fafc; }
  tr:hover { background: #eef2f7; }
  .badge { display: inline-block; padding: 3px 8px; border-radius: 6px; font-size: 11px; font-weight: 600; }
  .badge-room { background: #e0f2fe; color: #0369a1; }
  .badge-roll { background: #fef3c7; color: #92400e; font-family: monospace; font-size: 12px; }
  .badge-group { background: #ecfdf5; color: #065f46; }
</style>
</head>
<body>
<div class="header-card">
  <h1>Dhaka College &mdash; International Hall Student Database</h1>
  <p>Session 2025-2026 | Extracted from Official Admission Forms, Payment Receipts & Hall Directory</p>
</div>
<div class="stats-bar">
  <div class="stat-chip">Total Students: <span>28</span></div>
  <div class="stat-chip">Hostel: <span>International Hall</span></div>
  <div class="stat-chip">Science: <span>25</span></div>
  <div class="stat-chip">Humanities: <span>2</span></div>
  <div class="stat-chip">Business Studies: <span>1</span></div>
  <div class="stat-chip">Files Created: <span>.XLSX, .CSV, .HTML</span></div>
</div>
<div class="toolbar">
  <input type="text" id="searchInput" class="search-box" placeholder="Search by name, roll, phone, room, nid, district..." onkeyup="filterTable()">
  <button class="btn btn-success" onclick="copyTableForGoogleSheets()">📋 Copy All Data for Google Sheets</button>
  <button class="btn btn-primary" onclick="window.print()">🖨️ Print / Save PDF</button>
</div>
<div class="table-container">
<table id="studentsTable">
<thead>
<tr>
"""

for key, title, width in COLUMNS:
    html_content += f"  <th>{title}</th>\n"
html_content += "</tr>\n</thead>\n<tbody>\n"

for rec in records:
    html_content += "<tr>\n"
    for key, title, width in COLUMNS:
        v = rec.get(key, '')
        if key == 'hall_room':
            html_content += f'  <td><span class="badge badge-room">{v}</span></td>\n'
        elif key in ('short_roll', 'full_class_roll', 'ssc_roll'):
            html_content += f'  <td><span class="badge badge-roll">{v}</span></td>\n'
        elif key == 'group':
            html_content += f'  <td><span class="badge badge-group">{v}</span></td>\n'
        else:
            html_content += f"  <td>{v}</td>\n"
    html_content += "</tr>\n"

html_content += """</tbody>
</table>
</div>
<script>
function filterTable() {
  const input = document.getElementById('searchInput');
  const filter = input.value.toLowerCase();
  const table = document.getElementById('studentsTable');
  const tr = table.getElementsByTagName('tr');
  for (let i = 1; i < tr.length; i++) {
    const text = tr[i].textContent || tr[i].innerText;
    tr[i].style.display = text.toLowerCase().includes(filter) ? '' : 'none';
  }
}
function copyTableForGoogleSheets() {
  const table = document.getElementById('studentsTable');
  let tsv = '';
  for (let r = 0; r < table.rows.length; r++) {
    let row = [];
    for (let c = 0; c < table.rows[r].cells.length; c++) {
      let cellText = table.rows[r].cells[c].innerText.replace(/\\t/g, ' ').replace(/\\n/g, ' ').trim();
      row.push(cellText);
    }
    tsv += row.join('\\t') + '\\n';
  }
  navigator.clipboard.writeText(tsv).then(() => {
    alert('Copied all 28 students data to clipboard! Open Google Sheets, click on cell A1, and press Ctrl+V to paste.');
  });
}
</script>
</body>
</html>
"""

with open(html_path, 'w', encoding='utf-8') as f:
    f.write(html_content)

print(f"Saved HTML viewer to {html_path}")
print("ALL GENERATION COMPLETED SUCCESSFULLY!")
